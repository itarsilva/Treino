"""Gera a apresentação do AGM a partir do modelo (.pptx) e da planilha de dados (.xlsx).

Uso:
    python gerar_agm.py --modelo AGM_modelo.pptx --dados AGM_dados.xlsx --saida AGM_Setembro_2026.pptx

O modelo é a própria apresentação do AGM exportada do Google Slides. Os slides
são localizados pelos IDs dos objetos do Google Slides (estáveis enquanto o
mesmo arquivo for usado como modelo).
"""
import argparse
import copy
import os
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.table import _Cell
from pptx.util import Emu, Pt

import graficos as G
from analises import Analises, calc_nps, fnps, nps_zona
from dados import (Dados, div, fbrl, fbrl_curto, fdec, fint, fmil, fpct, nome_mes, num,
                   somar_mes)

# shape âncora (ID do Google Slides) que identifica cada slide do modelo
ANCORAS = {
    "capa": 163, "digital_fechados": 193, "digital_pipeline": 213, "digital_leads": 224,
    "digital_resumo": 245, "digital_status": 257, "proximos_digital": 270, "eventos": 316,
    "destaque": 333, "varejo_fechados": 347, "eventos_varejo": 375, "pipeline_eventos": 393,
    "pipeline_varejo": 407, "redes_alcance": 431, "redes_interacoes": 452, "redes_seguidores": 473,
    "concorrencia": 491, "pr": 506, "nps_relatorio": 532, "nps_projetos": 564, "nps_pontos": 595,
    "nps_sac": 607, "proximos_nps": 639,
}
ANALISE_Y, ANALISE_H = Emu(9330000), Emu(800000)


# ---------------------------------------------------------------- utilidades pptx

def shapes_por_id(slide):
    return {sh.shape_id: sh for sh in slide.shapes}


def _texto_par(p, texto):
    runs = p.runs
    if not runs:
        p.add_run()
        runs = p.runs
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def texto(sh, valor, par=None):
    """Troca o texto mantendo a formatação do primeiro run. `par` = índice do parágrafo."""
    ps = sh.text_frame.paragraphs
    if par is not None:
        _texto_par(ps[par], valor)
        return
    _texto_par(ps[0], valor)
    for p in ps[1:]:
        p._p.getparent().remove(p._p)


def linhas_texto(sh, linhas):
    """Um parágrafo por linha, copiando o estilo do primeiro parágrafo."""
    ps = sh.text_frame.paragraphs
    primeiro = copy.deepcopy(ps[0]._p)
    demais = copy.deepcopy(ps[min(1, len(ps) - 1)]._p)  # mantém o espaçamento entre itens
    corpo = ps[0]._p.getparent()
    for p in ps:
        corpo.remove(p._p)
    for i, _ in enumerate(linhas):
        corpo.append(copy.deepcopy(primeiro if i == 0 else demais))
    for p, ln in zip(sh.text_frame.paragraphs, linhas):
        _texto_par(p, ln)


def remover(sh):
    sh._element.getparent().remove(sh._element)


def novo_id(slide):
    return max(sh.shape_id for sh in slide.shapes) + 1


def clonar(slide, sh, dy=0):
    el = copy.deepcopy(sh._element)
    sh._element.getparent().append(el)
    el.xpath("./*[1]/p:cNvPr")[0].set("id", str(novo_id(slide)))
    novo = [s for s in slide.shapes if s._element is el][0]
    novo.top = sh.top + dy
    return novo


def trocar_cor(sh, cores_de, cor_para):
    for c in sh._element.iter(qn("a:srgbClr")):
        if c.get("val").upper() in cores_de:
            c.set("val", cor_para)


def trocar_imagem(slide, pic, png, top=None, height=None):
    novo = slide.shapes.add_picture(png, pic.left, top if top is not None else pic.top,
                                    pic.width, height if height is not None else pic.height)
    pic._element.addprevious(novo._element)
    rid = pic._element.blipFill.blip.rEmbed
    remover(pic)
    slide.part.drop_rel(rid)  # descarta a imagem antiga do arquivo
    return novo


def celula(tc, valor):
    c = _Cell(tc, None)
    _texto_par(c.text_frame.paragraphs[0], "" if valor is None else str(valor))
    for p in c.text_frame.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def tabela(sh, linhas, modelos=(1, 2), total=None, cabecalho=None):
    """Reescreve as linhas de uma tabela. `modelos` = índices das linhas usadas como estilo
    (alternadas); `total` = linha final com estilo da última linha do modelo."""
    tbl = sh.table._tbl
    trs = tbl.tr_lst
    estilos = [copy.deepcopy(trs[i]) for i in modelos]
    estilo_total = copy.deepcopy(trs[-1])
    if cabecalho:
        for tc, v in zip(trs[0].tc_lst, cabecalho):
            celula(tc, v)
    for tr in trs[1:]:
        tbl.remove(tr)
    for i, ln in enumerate(linhas):
        est = ln[0] if isinstance(ln, tuple) and isinstance(ln[0], int) else None
        vals = ln[1] if est is not None else ln
        tr = copy.deepcopy(estilos[est if est is not None else i % len(estilos)])
        for tc, v in zip(tr.tc_lst, vals):
            celula(tc, v)
        tbl.append(tr)
    if total:
        tr = copy.deepcopy(estilo_total)
        for tc, v in zip(tr.tc_lst, total):
            celula(tc, v)
        tbl.append(tr)
    sh.height = Emu(sum(int(tr.get("h")) for tr in tbl.tr_lst))


def remover_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
            return


def caixa_analise(slide, txt, largura=Emu(13700000)):
    if not txt:
        return
    tb = slide.shapes.add_textbox(Emu(1005840), ANALISE_Y, largura, ANALISE_H)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    for t, cor, negrito in (("Análise  ", "E100FF", True), (txt, "1F1F1F", False)):
        r = p.add_run()
        r.text = t
        r.font.size, r.font.bold, r.font.name = Pt(16), negrito, "Poppins"
        r.font.color.rgb = RGBColor.from_string(cor)
    notas(slide, txt)


def notas(slide, txt):
    if txt:
        slide.notes_slide.notes_text_frame.text = txt


# ---------------------------------------------------------------- gerador

class Gerador:
    def __init__(self, modelo, dados):
        self.prs = Presentation(modelo)
        self.d = Dados(dados)
        self.a = Analises(self.d)
        self.ref = self.d.ref
        self.tmp = tempfile.mkdtemp(prefix="agm_")
        self.slides = {}
        for s in self.prs.slides:
            ids = {sh.shape_id for sh in s.shapes}
            for k, anc in ANCORAS.items():
                if anc in ids:
                    self.slides[k] = s
        G.preparar_fontes()

    def png(self, nome):
        return os.path.join(self.tmp, nome + ".png")

    def run(self, saida):
        etapas = [self.capa, self.digital_fechados, self.digital_pipeline, self.digital_leads,
                  self.digital_resumo, self.digital_status, self.proximos_digital, self.eventos,
                  self.destaque, self.varejo, self.pipelines, self.redes, self.concorrencia, self.pr,
                  self.nps, self.proximos_nps]
        for etapa in etapas:
            etapa()
        self.prs.save(saida)
        return saida

    def _s(self, chave):
        s = self.slides.get(chave)
        if s is None:
            raise SystemExit(f"Slide '{chave}' não encontrado no modelo.")
        return s, shapes_por_id(s)

    # --- capa
    def capa(self):
        s, sh = self._s("capa")
        texto(sh[163], f"{nome_mes(self.ref, cap=True)}/{self.ref[:4]}", par=2)
        if 170 in sh:  # faixa vermelha (placeholder) do modelo
            remover(sh[170])

    # --- Digital
    def _meses_ate_ref(self, aba, campo, desde_primeiro=False):
        meses = self.d.meses_ano()
        if desde_primeiro:
            com = [m for m in meses if (self.d.valor(aba, m, campo) or 0) > 0]
            if com:
                meses = meses[meses.index(com[0]):]
        return meses

    def digital_fechados(self):
        s, sh = self._s("digital_fechados")
        d = self.d
        meses = self._meses_ate_ref("Digital", "valor_fechado", desde_primeiro=True)
        vals = [d.valor("Digital", m, "valor_fechado") or 0 for m in meses]
        pic = sh[190]
        G.barras(self.png("dig_fech"), [nome_mes(m) for m in meses], vals, pic.width, pic.height,
                 fmt=lambda v: fbrl_curto(v))
        trocar_imagem(s, pic, self.png("dig_fech"))
        receita = sum(d.valor("Digital", m, "valor_fechado") or 0 for m in d.meses_ano())
        invest = sum(d.valor("Digital", m, "investimento") or 0 for m in d.meses_ano())
        roas = fpct(div(receita, invest) * 100, 0) if invest else "—"
        texto(sh[193], fbrl(receita))
        texto(sh[196], fbrl(invest))
        texto(sh[199], roas)
        caixa_analise(s, self.a.get("digital_fechados", receita, roas))

    def _pipeline(self, chave, origem, pic_id, total_id, cor):
        s, sh = self._s(chave)
        rs = self.d.linhas("Pipeline", self.ref, origem=origem)
        cats = [str(r["fase"]) for r in rs]
        vals = [num(r["valor"]) or 0 for r in rs]
        pic = sh[pic_id]
        G.barras_h(self.png(chave), cats, vals, pic.width, pic.height, cor=cor, fmt=fbrl)
        trocar_imagem(s, pic, self.png(chave))
        texto(sh[total_id], fbrl(sum(vals)))
        caixa_analise(s, self.a.get(chave))

    def digital_pipeline(self):
        self._pipeline("digital_pipeline", "Digital", 210, 213, G.ROXO)

    def pipelines(self):
        self._pipeline("pipeline_eventos", "Eventos", 390, 393, G.CIANO)
        self._pipeline("pipeline_varejo", "Varejo", 404, 407, G.ROXO)

    def digital_leads(self):
        s, sh = self._s("digital_leads")
        meses = self.d.meses_ano()
        leads = [self.d.valor("Digital", m, "leads") for m in meses]
        inv = [self.d.valor("Digital", m, "investimento") for m in meses]
        pic = sh[224]
        topo, fim = pic.top, sh[226].top + sh[226].height
        G.linha(self.png("leads"), [nome_mes(m, abrev=True) for m in meses], leads, pic.width,
                fim - topo, fmt=fint,
                linha_extra=("Investimento (R$)", [f"{fdec(v / 1000, 1)}k" if v else "—" for v in inv]))
        for i in range(225, 236):
            remover(sh[i])
        trocar_imagem(s, pic, self.png("leads"), height=fim - topo)
        caixa_analise(s, self.a.get("digital_leads"))

    def digital_resumo(self):
        s, sh = self._s("digital_resumo")
        d = self.d
        if d.parcial:
            meses = [somar_mes(self.ref, -i) for i in (3, 2, 1, 0)] + [d.parcial]
        else:
            meses = [somar_mes(self.ref, -i) for i in (4, 3, 2, 1, 0)]
        g = lambda m, c: d.valor("Digital", m, c)
        pct = lambda a, b: fpct(div(a, b) * 100) if div(a, b) is not None else "—"
        linhas = [
            ["Visitas"] + [fint(g(m, "visitas")) for m in meses],
            ["Leads"] + [fint(g(m, "leads")) for m in meses],
            ["Leads qualificados"] + [fint(g(m, "leads_qualificados")) for m in meses],
            ["Negócios fechados"] + [fint(g(m, "negocios_fechados")) for m in meses],
            ["Valor"] + [fbrl(g(m, "valor_fechado")) for m in meses],
            ["Visitas > Leads"] + [pct(g(m, "leads"), g(m, "visitas")) for m in meses],
            ["Leads > Qualificados"] + [pct(g(m, "leads_qualificados"), g(m, "leads")) for m in meses],
            ["Qualificados > Vendas"] + [pct(g(m, "negocios_fechados"), g(m, "leads_qualificados")) for m in meses],
            ["Total (Vendas/Leads)"] + [pct(g(m, "negocios_fechados"), g(m, "leads")) for m in meses],
            ["Orçamento (R$)"] + [fbrl(g(m, "investimento")) for m in meses],
        ]
        cab = ["Métrica"] + [nome_mes(m, cap=True) + ("*" if m == d.parcial else "") for m in meses]
        tabela(sh[245], linhas, modelos=(1, 2), cabecalho=cab)
        if d.parcial:
            texto(sh[246], f"*Dados de {nome_mes(d.parcial)} até {d.parcial_ate}.")
        else:
            remover(sh[246])
        caixa_analise(s, self.a.get("digital_resumo"))

    def digital_status(self):
        s, sh = self._s("digital_status")
        d = self.d
        rs = d.linhas("Leads_Status", self.ref)
        pic = sh[256]
        G.barras_h(self.png("status"), [str(r["status"]) for r in rs], [num(r["quantidade"]) for r in rs],
                   pic.width, pic.height, cor=G.ROXO, fmt=fint)
        trocar_imagem(s, pic, self.png("status"))

        def sdr(tab, mes, rotulo):
            rs = d.linhas("Leads_SDR", mes)
            vals = [[str(r["sdr"]), fint(num(r["leads"]) or 0)] for r in rs]
            tabela(tab, vals, modelos=(1, 2), cabecalho=[rotulo, "Leads"],
                   total=["Total", fint(sum(num(r["leads"]) or 0 for r in rs))])

        sdr(sh[257], self.ref, nome_mes(self.ref, cap=True))
        if d.parcial and d.linhas("Leads_SDR", d.parcial):
            sdr(sh[258], d.parcial, nome_mes(d.parcial, cap=True) + "*")
        else:
            remover(sh[258])
        caixa_analise(s, self.a.get("digital_status"), largura=Emu(8500000))

    def _lista_numerada(self, chave, grupos, itens, preencher):
        s, sh = self._s(chave)
        grupos = [[sh[i] for i in g] for g in grupos]
        passo = grupos[1][0].top - grupos[0][0].top
        while len(grupos) < len(itens):
            ultimo = grupos[-1]
            grupos.append([clonar(s, x, passo) for x in ultimo])
        if len(itens) > 3:  # aperta o espaçamento para caber até o rodapé
            y0 = grupos[0][0].top
            novo = min(passo, int((8900000 - y0) / len(itens)))
            for i, g in enumerate(grupos):
                dy = y0 + novo * i - g[0].top
                for x in g:
                    x.top = x.top + dy
        for g in grupos[len(itens):]:
            for x in g:
                remover(x)
        for i, (g, item) in enumerate(zip(grupos, itens), 1):
            texto(g[0], f"{i:02d}")
            preencher(g, item)

    def proximos_digital(self):
        itens = [str(r["titulo"] or r["texto"]) for r in self.d.textos("proximos_digital")]
        if not itens:
            return
        self._lista_numerada("proximos_digital", [(269, 270), (271, 272), (273, 274)], itens,
                             lambda g, it: texto(g[1], it))

    def proximos_nps(self):
        itens = [(str(r["titulo"] or ""), str(r["texto"] or "")) for r in self.d.textos("proximos_nps")]
        if not itens:
            return
        self._lista_numerada("proximos_nps", [(637, 638, 639), (640, 641, 642), (643, 644, 645)], itens,
                             lambda g, it: (texto(g[1], it[0]), texto(g[2], it[1])))

    # --- Eventos
    def eventos(self):
        s, sh = self._s("eventos")
        rs = self.d.linhas("Eventos", self.ref)
        texto(sh[293], f"Eventos {nome_mes(self.ref, cap=True)}")
        grupos = [(294, 295, 296, 297), (298, 299, 300), (301, 302, 303), (304, 305, 306, 307),
                  (308, 309, 310, 311), (312, 313, 314, 315)]
        cores = []
        for g in grupos:
            for c in sh[g[0]]._element.iter(qn("a:srgbClr")):
                cores.append(c.get("val"))
                break
        modelo = [sh[i] for i in grupos[0]]
        y0 = modelo[0].top
        passo = 1097280 if len(rs) <= 6 else int((8700000 - y0) / max(len(rs) - 1, 1))
        for i, r in enumerate(rs):
            novos = [clonar(s, x, passo * i) for x in modelo]
            trocar_cor(novos[0], {"00E1F0", "7D7BFF", "E100FF"}, cores[i % len(cores)])
            texto(novos[1], str(r.get("data") or ""))
            texto(novos[2], str(r.get("evento") or ""))
            if r.get("local"):
                texto(novos[3], str(r["local"]))
            else:
                remover(novos[3])
        for g in grupos:
            for i in g:
                remover(sh[i])
        com_leads = [r for r in rs if num(r.get("leads_total")) is not None]
        if com_leads:
            tabela(sh[316], [[str(r["evento"]), fint(num(r["leads_total"])), fint(num(r.get("leads_novos")))]
                             for r in com_leads], modelos=(1, 2))
        else:
            remover(sh[316])
        caixa_analise(s, self.a.get("eventos"))

    def destaque(self):
        s, sh = self._s("destaque")
        tit = self.d.textos("destaque_titulo")
        if not tit:
            remover_slide(self.prs, s)
            return
        texto(sh[325], str(tit[0]["titulo"] or tit[0]["texto"]))
        itens = [str(r["titulo"] or r["texto"]) for r in self.d.textos("destaque_item")]
        if itens:
            linhas_texto(sh[326], [f"{i}.  {t}" for i, t in enumerate(itens, 1)])
        else:
            remover(sh[326])
        kpis = self.d.textos("destaque_kpi")
        cartoes = [(327, 328, 329), (330, 331, 332)]
        for i, ids in enumerate(cartoes):
            if i < len(kpis):
                texto(sh[ids[1]], str(kpis[i]["titulo"] or ""))
                texto(sh[ids[2]], str(kpis[i]["texto"] or ""))
            else:
                for j in ids:
                    remover(sh[j])
        tab = self.d.textos("destaque_tabela")
        if tab:
            qtds = [num(r["texto"]) for r in tab]
            total = ["Total geral", fint(sum(q for q in qtds if q))] if all(q is not None for q in qtds) else None
            tabela(sh[333], [[str(r["titulo"]), str(r["texto"])] for r in tab], modelos=(1, 2), total=total)
        else:
            remover(sh[333])

    # --- Varejo / Eventos (receita)
    def varejo(self):
        d = self.d
        meses = d.meses_ano()
        varejo = [d.valor("Receita", m, "varejo") or 0 for m in meses]
        ev = [d.valor("Receita", m, "eventos") or 0 for m in meses]
        inv = sum(d.valor("Receita", m, "investimento_eventos") or 0 for m in meses)
        ytd_v, ytd_e = sum(varejo), sum(ev)
        v70 = ytd_v * 0.7
        total = ytd_e + v70
        roas = fpct(div(total, inv) * 100, 0) if inv else "—"

        s, sh = self._s("varejo_fechados")
        pic = sh[344]
        G.barras(self.png("varejo"), [nome_mes(m, abrev=True) for m in meses], varejo, pic.width, pic.height,
                 fmt=lambda v: f"R$ {fdec(v / 1e6, 2)} mi")
        trocar_imagem(s, pic, self.png("varejo"))
        texto(sh[347], fbrl(ytd_v))
        texto(sh[350], fbrl(v70))
        caixa_analise(s, self.a.get("varejo_fechados", ytd_v))

        s, sh = self._s("eventos_varejo")
        texto(sh[363], fbrl(inv) if inv else "—")
        texto(sh[366], fbrl(ytd_e))
        texto(sh[369], fbrl(v70))
        texto(sh[372], fbrl(total))
        texto(sh[375], roas)
        mev = [(m, v) for m, v in zip(meses, ev) if v]
        pic = sh[377]
        G.barras(self.png("ev"), [nome_mes(m, abrev=True) for m, _ in mev], [v for _, v in mev],
                 pic.width, pic.height, cor=G.CIANO, rotulo=16, eixo=19,
                 fmt=lambda v: f"{fint(v / 1e3)} mil")
        trocar_imagem(s, pic, self.png("ev"))
        pic = sh[379]
        G.barras(self.png("ev_var"), [nome_mes(m, abrev=True) for m in meses], varejo, pic.width, pic.height,
                 cor=G.ROXO, rotulo=16, eixo=19, fmt=lambda v: f"{fdec(v / 1e6, 1)} mi")
        trocar_imagem(s, pic, self.png("ev_var"))
        notas(s, self.a.get("eventos_varejo", total, inv, roas))

    # --- Redes
    def redes(self):
        meses = [somar_mes(self.ref, -2), somar_mes(self.ref, -1), self.ref]
        ordem = ["Instagram", "Facebook", "TikTok", "YouTube", "LinkedIn"]
        presentes = {str(r["rede"]) for r in self.d.abas["Redes"] if r["mes"] in meses}
        redes = [r for r in ordem if r in presentes] + sorted(presentes - set(ordem))
        cfg = [("redes_alcance", "alcance", 428, (430, 431, 433, 434, 436, 437), 438, 427, "do alcance"),
               ("redes_interacoes", "interacoes", 449, (451, 452, 454, 455, 457, 458), 459, 448, "das interações"),
               ("redes_seguidores", "novos_seguidores", 470, (472, 473, 475, 476, 478, 479), 480, 469,
                "dos novos seguidores")]
        for chave, campo, pic_id, caixas, comentario, subt, rotulo in cfg:
            s, sh = self._s(chave)
            series = []
            for m, cor in zip(meses, (G.CIANO, G.ROXO, G.MAGENTA)):
                vals = [self.d.valor("Redes", m, campo, rede=r) for r in redes]
                series.append((nome_mes(m, abrev=True, cap=True), vals, cor))
            pic = sh[pic_id]
            G.agrupado(self.png(chave), redes, series, pic.width, pic.height, fmt=fint)
            trocar_imagem(s, pic, self.png(chave))
            texto(sh[subt], " × ".join(nome_mes(m, cap=True) for m in meses))
            for k, m in enumerate(meses):
                texto(sh[caixas[2 * k]], f"Total {nome_mes(m)}")
                texto(sh[caixas[2 * k + 1]], fint(self.d.valor("Redes", m, campo)))
            manual = self.d.analise_manual(chave)
            caixa, txt = self.a.redes(campo, meses, rotulo)
            texto(sh[comentario], caixa)
            caixa_analise(s, manual or txt)

    def concorrencia(self):
        s, sh = self._s("concorrencia")
        rs = self.d.linhas("Concorrencia", self.ref)
        rs.sort(key=lambda r: (0 if "ledwave" in str(r["conta"]).lower() else 1,
                               -(num(r.get("novos_seguidores")) or 0)))
        linhas = []
        for i, r in enumerate(rs):
            nosso = "ledwave" in str(r["conta"]).lower()
            vals = [str(r["conta"]), fmil(num(r.get("seguidores"))),
                    f"↑ {fmil(num(r.get('novos_seguidores')))}", fint(num(r.get("publicacoes")))]
            linhas.append((0 if nosso else 1 + (i % 2), vals))
        tabela(sh[491], linhas, modelos=(1, 2, 3))
        caixa_analise(s, self.a.get("concorrencia"))

    def pr(self):
        s, sh = self._s("pr")
        texto(sh[503], fint(self.d.valor("PR", self.ref, "materias")))
        texto(sh[506], fbrl(self.d.valor("PR", self.ref, "earned_media")))
        dest = [str(r["titulo"] or r["texto"]) for r in self.d.textos("pr_destaque")]
        if dest:
            linhas_texto(sh[508], [f"•  {t}" for t in dest])
        else:
            remover(sh[507])
            remover(sh[508])
        veic = self.d.textos("pr_veiculos")
        if veic:
            texto(sh[509], f"Veículos: {veic[0]['texto'] or veic[0]['titulo']}")
        else:
            remover(sh[509])
        notas(s, self.a.get("pr"))

    # --- NPS
    def _nps_slide(self, chave, ids, p, pa, dt, env, com):
        s, sh = self._s(chave)
        ring, val, zona, cp, cpa, cd, cr, ce, ct, obs = ids
        resp = p + pa + dt
        n = calc_nps(p, pa, dt)
        z = nps_zona(n)
        cor = "E100FF" if n is not None and n >= 75 else "7D7BFF"
        texto(sh[val], fnps(n) if n is not None else "—")
        texto(sh[zona], z.upper())
        trocar_cor(sh[ring], {"E100FF", "7D7BFF"}, cor)
        trocar_cor(sh[zona], {"E100FF", "7D7BFF"}, cor)
        pc = lambda x: f"{fpct(x / resp * 100, 1).replace(',0%', '%') if resp else '0,0%'} ({fint(x)})"
        pc0 = lambda x: pc(x) if x else f"0,0% ({fint(x)})"
        texto(sh[cp], pc0(p))
        texto(sh[cpa], pc0(pa))
        texto(sh[cd], pc0(dt))
        texto(sh[cr], fint(resp))
        texto(sh[ce], fint(env))
        texto(sh[ct], fpct(div(resp, env) * 100, 1) if env else "—")
        texto(sh[obs], f"{fint(p)} promotores, {fint(pa)} passivos e {fint(dt)} detratores · "
                       f"{fint(com)} {'respostas' if com != 1 else 'resposta'} com comentário.")
        caixa_analise(s, self.a.get(chave, chave, p, pa, dt, env))

    def nps(self):
        d = self.d
        areas = {}
        for area in ("Projetos", "SAC"):
            areas[area] = [d.valor("NPS", self.ref, c, area=area) or 0
                           for c in ("promotores", "passivos", "detratores", "enviadas", "comentarios")]
        tot = [a + b for a, b in zip(areas["Projetos"], areas["SAC"])]
        ids = {
            "nps_relatorio": (530, 532, 533, 536, 539, 542, 545, 548, 551, 552),
            "nps_projetos": (562, 564, 565, 568, 571, 574, 577, 580, 583, 584),
            "nps_sac": (605, 607, 608, 611, 614, 617, 620, 623, 626, 627),
        }
        self._nps_slide("nps_relatorio", ids["nps_relatorio"], *tot)
        self._nps_slide("nps_projetos", ids["nps_projetos"], *areas["Projetos"])
        self._nps_slide("nps_sac", ids["nps_sac"], *areas["SAC"])

        s, sh = self._s("nps_pontos")
        rs = d.linhas("NPS_Pontos", self.ref)
        linhas, resumo = [], []
        base = max([num(r.get("base")) or 0 for r in rs] + [0])
        for r in rs:
            resp, b = num(r.get("respostas")) or 0, num(r.get("base")) or base
            n = calc_nps(num(r.get("promotores")), num(r.get("passivos")), num(r.get("detratores")))
            resumo.append((str(r["ponto"]), n))
            linhas.append([str(r["ponto"]), f"{fpct(div(resp, b) * 100, 1).replace(',0%', '%')} / {fint(resp)}",
                           fnps(n), nps_zona(n)])
        tabela(sh[595], linhas, modelos=(1, 2))
        texto(sh[594], f"{fint(areas['Projetos'][3])} enviadas · {fint(base)} respondidas")
        caixa_analise(s, self.a.get("nps_pontos", resumo))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--modelo", required=True)
    ap.add_argument("--dados", required=True)
    ap.add_argument("--saida")
    a = ap.parse_args()
    g = Gerador(a.modelo, a.dados)
    saida = a.saida or f"AGM_{nome_mes(g.ref, cap=True)}_{g.ref[:4]}.pptx"
    print("Gerado:", g.run(saida))


if __name__ == "__main__":
    main()
