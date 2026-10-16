"""Cria a planilha modelo de dados do AGM (.xlsx, pronta para subir no Google Sheets).

Uso:
    python criar_planilha.py saida.xlsx [--dados dados.json]

`--dados` (opcional) preenche as abas com linhas já existentes, no formato
{"Digital": [{"mes": "2025-10", "leads": 724, ...}, ...], "Config": {...}}.
"""
import argparse
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from esquema import ABAS, CHAVES_ANALISE, CONFIG_PADRAO, SECOES_TEXTO

PRETO = PatternFill("solid", fgColor="000000")
LILAS = PatternFill("solid", fgColor="ECE6FF")
CAB = Font(name="Poppins", bold=True, color="FFFFFF")
NEGRITO = Font(name="Poppins", bold=True)
NORMAL = Font(name="Poppins")


def leia_me(ws):
    linhas = [
        ("AGM Marketing — planilha de dados", NEGRITO),
        ("", None),
        ("Como usar todo mês", NEGRITO),
        ("1. Na aba Config, troque mes_referencia para o mês de resultados (AAAA-MM). Ex.: apresentação de outubro → 2026-09.", None),
        ("2. Em cada aba, ADICIONE novas linhas para o mês (não apague os meses anteriores — eles alimentam os gráficos e comparações).", None),
        ("3. Se houver dados parciais do mês atual (ex.: até 07/10), lance-os com o mês seguinte e preencha parcial_ate na Config.", None),
        ("4. Valores em R$ podem ser digitados como número (ex.: 15623,49). Percentuais e totais são calculados automaticamente.", None),
        ("5. Avise o Claude: ele lê a planilha, gera a apresentação e escreve as análises.", None),
        ("", None),
        ("Seções da aba Textos", NEGRITO),
        *[(f"• {s} — {d}", None) for s, d in SECOES_TEXTO],
        ("", None),
        ("Chaves para a seção 'analise'", NEGRITO),
        (", ".join(CHAVES_ANALISE), None),
    ]
    for i, (txt, fonte) in enumerate(linhas, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = fonte or NORMAL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140


def criar(saida, dados=None):
    dados = dados or {}
    wb = Workbook()
    leia_me(wb.active)
    wb.active.title = "LEIA-ME"

    for aba, colunas in ABAS.items():
        ws = wb.create_sheet(aba)
        for j, (col, desc) in enumerate(colunas, 1):
            c = ws.cell(row=1, column=j, value=col)
            c.font, c.fill = CAB, PRETO
            d = ws.cell(row=2, column=j, value=desc)
            d.font, d.fill = Font(name="Poppins", italic=True, size=9, color="6B6B6B"), LILAS
            d.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[get_column_letter(j)].width = max(16, min(45, len(desc) * 0.6))
        ws.freeze_panes = "A3"

        if aba == "Config":
            cfg = {k: v for k, v, _ in CONFIG_PADRAO}
            cfg.update(dados.get("Config", {}))
            for k, _, desc in CONFIG_PADRAO:
                ws.append([k, cfg.get(k, ""), desc])
            ws.column_dimensions["C"].width = 70
            continue

        nomes = [c for c, _ in colunas]
        for linha in dados.get(aba, []):
            ws.append([linha.get(n) for n in nomes])

        # mês como texto, para o Sheets não converter em data
        for row in ws.iter_rows(min_row=3, max_col=1):
            for c in row:
                c.number_format = "@"
    wb.save(saida)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("saida")
    ap.add_argument("--dados")
    a = ap.parse_args()
    criar(a.saida, json.load(open(a.dados, encoding="utf-8")) if a.dados else None)
