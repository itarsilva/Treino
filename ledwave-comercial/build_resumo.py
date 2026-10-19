import sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

from conteudo import NICHOS, OBJECOES_GERAIS, STATUS_IDEIA, TIPOS_IDEIA
from conteudo2 import SPIN, ORCAMENTO
from conteudo3 import PROPOSITOS, DIFERENCIAIS, OBJECOES_COMERCIAL, CASES, NICHO_MAPA

OUT = sys.argv[1]
NAVY, MAGENTA, EDIT, WHITE, GREY = "3D3D3D", "8A5CF6", "FFF8D6", "FFFFFF", "F5F5F7"
F = "Arial"
thin = Side(style="thin", color="D0CCE0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
wb.remove(wb.active)


def sheet(name, tab=NAVY):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    return ws


def title(ws, text, sub, span):
    ws["A1"], ws["A2"] = text, sub
    ws["A1"].font = Font(name=F, size=18, bold=True, color=WHITE)
    ws["A2"].font = Font(name=F, size=10, italic=True, color="D9D2F0")
    for r in (1, 2):
        for c in range(1, span + 1):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20


def header(ws, row, cols, widths=None, color=MAGENTA):
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=F, size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
    ws.row_dimensions[row].height = 32
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def est(texts, widths):
    lines = 1
    for t, w in zip(texts, widths):
        if t is None:
            continue
        chars = max(int(w * 1.15), 8)
        lines = max(lines, sum(max(1, -(-len(p) // chars)) for p in str(t).split("\n")))
    return min(max(lines * 14 + 6, 20), 400)


def row(ws, r, values, widths, bold_first=True, edit_cols=(), zebra=False):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=F, size=10, bold=(bold_first and i == 1))
        c.alignment = WRAP
        c.border = BORDER
        if i in edit_cols:
            c.fill = PatternFill("solid", fgColor=EDIT)
        elif zebra:
            c.fill = PatternFill("solid", fgColor=GREY)
    ws.row_dimensions[r].height = est(values, widths)


def section(ws, r, text):
    ws.cell(row=r, column=1, value=text).font = Font(name=F, size=12, bold=True, color=NAVY)


N = len(NICHOS)
DIF = {d[0]: d for d in DIFERENCIAIS}

# ------------------------------------------------------------------ Início
ws = sheet("Início", MAGENTA)
title(ws, "Playbook Comercial LedWave · versão resumida", "Vender primeiro a LedWave, depois o porquê do cliente · apresentações de 10 a 20 min · v3 (out/2026)", 5)
w = [26, 70, 45, 16, 26]
for i, x in enumerate(w, start=1):
    ws.column_dimensions[get_column_letter(i)].width = x
passos = [
    ("1. Antes da reunião", "Aba Descoberta: perguntas SPIN e o que levantar para orçar o nicho."),
    ("2. Monte o roteiro", "Aba Gerador: escolha nicho e duração; o roteiro sai pronto em 7 blocos."),
    ("3. Descubra o propósito", "Aba Mapa de Venda: identifique por que este cliente compra (mídia, branding, experiência...) e use os diferenciais, resultados e cases daquele propósito."),
    ("4. Trate as objeções", "Aba Objeções: resposta curta + a prova que você leva."),
    ("5. Contribua", "Aba Ideias: uma linha por ideia. As aprovadas sobem para a matriz Nichos."),
]
section(ws, 4, "Como usar em 5 passos")
for i, (a, b) in enumerate(passos):
    row(ws, 5 + i, [a, b], w)
r = 6 + len(passos)
section(ws, r, "LedWave em 1 minuto (usado no bloco 3 do Gerador)")
header(ws, r + 1, ["Prova", "Como falar", "Fonte", "Status", "Observação"], color=NAVY)
PROVAS = [
    ("Origem", "Começamos com R$ 12 mil e uma ideia trazida de Xangai. Hoje somos referência em LED no Brasil.", "Exame", "OK"),
    ("Liderança", "Quatro anos seguidos entre as empresas que mais crescem no Brasil e nº 1 do segmento de painéis de LED no ranking Exame 2025.", "Ranking Exame Negócios em Expansão", "OK"),
    ("Projeto-ícone", "Fizemos o Boulevard.Live: 1.125 m², o maior painel de LED da América Latina.", "Blog LedWave; Inviron", "OK"),
    ("Presença", "Fábrica em Goiânia e unidades em São Paulo, Brasília e Orlando (EUA).", "Site LedWave", "OK"),
    ("Internacional", "Única empresa brasileira de LED com estande próprio na NRF 2026, em Nova York.", "Imprensa", "OK"),
    ("Clientes", "Nordestão, Rosewood, Beto Carrero World e Museu do Amanhã escolheram a LedWave.", "Time comercial", "Confirmar autorização"),
]
for i, p in enumerate(PROVAS):
    row(ws, r + 2 + i, list(p) + [""], w, edit_cols=(4, 5))
PROVA_ROW = r + 2  # first proof row
r = r + 3 + len(PROVAS)
section(ws, r, "Legenda")
c = ws.cell(row=r + 1, column=1, value="Campo para preencher")
c.fill, c.font, c.border = PatternFill("solid", fgColor=EDIT), Font(name=F, size=10), BORDER
ws.cell(row=r + 2, column=1, value="Versão completa (roleplay, técnicas, frases, prospecção detalhada): arquivo LedWave_Playbook_Comercial_Nichos.xlsx").font = Font(name=F, size=9, italic=True)

# ------------------------------------------------------------------ Gerador
wg = sheet("Gerador", MAGENTA)
title(wg, "Gerador de Roteiro", "Escolha nicho e duração nas células amarelas. Ajuste os percentuais da coluna B se quiser (devem somar 100%).", 6)
w = [28, 10, 10, 11, 30, 85]
for i, x in enumerate(w, start=1):
    wg.column_dimensions[get_column_letter(i)].width = x
for a, lbl in (("A4", "Nicho"), ("A5", "Duração (min)")):
    wg[a] = lbl
    wg[a].font = Font(name=F, size=11, bold=True, color=NAVY)
wg["B4"], wg["B5"] = NICHOS[1]["nicho"], 15
for a in ("B4", "B5"):
    wg[a].fill = PatternFill("solid", fgColor=EDIT)
    wg[a].font = Font(name=F, size=11, bold=True)
    wg[a].border = BORDER
wg.merge_cells("B4:E4")
dv = DataValidation(type="list", formula1="=Listas!$A$2:$A$60")
wg.add_data_validation(dv); dv.add("B4")
dv = DataValidation(type="list", formula1='"10,15,20"')
wg.add_data_validation(dv); dv.add("B5")
M = "MATCH($B$4,Nichos!$A:$A,0)"
col = lambda c: f"INDEX(Nichos!${c}:${c},{M})"
spin = lambda c: f"INDEX(Descoberta!${c}:${c},MATCH($B$4,Descoberta!$A:$A,0))"
prova = lambda i: f"INDEX('Início'!$B:$B,{PROVA_ROW + i})"
wg["F4"] = f'=IFERROR("Decisor: "&{col("B")}&"   ·   Propósitos de compra: "&{col("C")},"Escolha um nicho da lista")'
wg["F4"].font = Font(name=F, size=10, italic=True, color=MAGENTA)
wg["F4"].alignment = WRAP
wg.row_dimensions[4].height = 42
NL = "&CHAR(10)&"
blocos = [
    ("1. Abertura: pergunta", 0.08, "Pergunta provocativa + pausa de 3 s",
     f'"PERGUNTE: "&{col("E")}'),
    ("2. Dor e custo de esperar", 0.20, "SPIN + escuta 70/30",
     f'{col("D")}{NL}"SITUAÇÃO: "&{spin("B")}{NL}"PROBLEMA: "&{spin("C")}{NL}"IMPLICAÇÃO: "&{spin("D")}'),
    ("3. LedWave: o guia", 0.18, "StoryBrand + regra de 3",
     f'{prova(0)}{NL}{prova(1)}{NL}{prova(2)}{NL}"CASE: "&IF({col("J")}="","Boulevard.Live",{col("J")}){NL}"Ponte: E é por isso que estamos aqui."'),
    ("4. O porquê", 0.12, "Golden Circle + analogia",
     f'"CONFIRME O PROPÓSITO DE COMPRA ("&{col("C")}&") e use a aba Mapa de Venda."{NL}{col("F")}{NL}"FRASE (pausa): "&{col("H")}'),
    ("5. Três ganhos e solução", 0.20, "Regra de 3 + antes/depois",
     f'{col("G")}{NL}"SOLUÇÃO E KPIs: "&{col("I")}'),
    ("6. Objeções antecipadas", 0.10, "Ouvir, acolher, explorar, responder com prova",
     f'"MAIS PROVÁVEIS: "&{col("K")}&" (respostas na aba Objeções)"'),
    ("7. Fechamento", 0.12, "Peak-End + escolha entre duas datas",
     f'"NECESSIDADE: "&{spin("E")}{NL}"FRASE: "&{col("H")}{NL}{col("L")}'),
]
header(wg, 7, ["Bloco", "% do tempo", "Minutos", "Termina em", "Técnica", "Fala do nicho"])
for i, (b, pct, tec, fala) in enumerate(blocos):
    r = 8 + i
    vals = [b, pct, f"=B{r}*$B$5", f"=SUM($C$8:C{r})", tec, f'=IFERROR({fala},"Escolha um nicho da lista")']
    for j, v in enumerate(vals, start=1):
        c = wg.cell(row=r, column=j, value=v)
        c.font = Font(name=F, size=10, bold=(j == 1))
        c.alignment = WRAP
        c.border = BORDER
    wg.cell(row=r, column=2).number_format = "0%"
    wg.cell(row=r, column=2).fill = PatternFill("solid", fgColor=EDIT)
    for j in (3, 4):
        wg.cell(row=r, column=j).number_format = "0.0"
    wg.row_dimensions[r].height = 150 if i in (1, 2, 3, 4, 6) else 75
r = 8 + len(blocos)
wg.cell(row=r, column=1, value="Total").font = Font(name=F, size=10, bold=True)
wg.cell(row=r, column=2, value=f"=SUM(B8:B{r-1})").number_format = "0%"
wg.cell(row=r, column=3, value=f"=SUM(C8:C{r-1})").number_format = "0.0"
wg.cell(row=r, column=5, value=f'=IF(ROUND(B{r},2)=1,"OK","Ajuste para somar 100%")').font = Font(name=F, size=10, bold=True, color=MAGENTA)
wg.freeze_panes = "A8"

# ------------------------------------------------------------------ Nichos
wn = sheet("Nichos", MAGENTA)
title(wn, "Nichos: o porquê de cada cliente", "Uma linha por nicho, fonte do Gerador. Mantenha a ordem das colunas; melhore os textos à vontade.", 13)
cols = ["Nicho", "Decisor típico", "Propósitos de compra", "Dor e custo de esperar", "Pergunta de abertura",
        "O porquê + história", "Três ganhos", "Frase de efeito", "Solução e KPIs", "Cases", "Objeções prováveis",
        "Fechamento e próximo passo", "Ideias no banco"]
w = [24, 24, 22, 45, 38, 50, 52, 26, 45, 20, 26, 45, 10]
header(wn, 3, cols, w)
for i, n in enumerate(NICHOS):
    r = 4 + i
    pr, cs, ob = NICHO_MAPA[n["nicho"]]
    vals = [
        n["nicho"], n["decisor"], " · ".join(pr),
        f"{n['dor']}\nCusto de esperar: {n['inacao']}",
        n["gancho"],
        f"{n['proposito']}\nHistória: {n['historia']}",
        f"Receita: {n['arg_receita']}\nExperiência: {n['arg_exp']}\nOperação: {n['arg_oper']}",
        n["frase"],
        f"{n['solucao']}\nKPIs: {n['kpi']}",
        "; ".join(cs), "; ".join(ob),
        f"{n['fechamento']}\nPróximo passo: {n['cta']}",
        f"=COUNTIF(Ideias!$C:$C,A{r})",
    ]
    row(wn, r, vals, w[:12], zebra=(i % 2 == 1))
    wn.cell(row=r, column=8).font = Font(name=F, size=10, bold=True, color=MAGENTA)
    wn.cell(row=r, column=13).alignment = Alignment(horizontal="center", vertical="top")
wn.freeze_panes = "B4"
wn.auto_filter.ref = f"A3:M{3+N}"

# ------------------------------------------------------------------ Mapa de Venda
MAPA_EXTRA = {
    "Mídia": (["Veículo de mídia", "Software e programática", "Joint venture", "Modelo comercial"], ["Política local / licenças", "Crédito", "SLA"], "Painel apagado é aluguel perdido."),
    "Eficiência operacional": (["Software e programática", "Solução 360", "Capilaridade de atendimento"], ["Fobia digital", "Gestão de conteúdo", "Percepção de valor"], "Uma marca, mil lojas, um clique."),
    "Branding": (["Customização", "Liderança", "Prova social"], ["Preço", "Percepção de valor"], "Toda empresa comunica. Poucas encantam."),
    "Retrofit / arquitetura": (["Customização", "Única fábrica", "Certificações", "Garantias"], ["Fabricação de estrutura metálica", "Política local / licenças", "Consumo de energia"], "O prédio vira mensagem."),
    "Retail media": (["Software e programática", "Modelo comercial", "Joint venture"], ["Gestão de conteúdo", "Aprovação por financeiro", "Cenário mercadológico"], "Sua loja já tem audiência. Falta transformar audiência em receita."),
    "Trade": (["Software e programática", "Capilaridade de atendimento", "Solução 360"], ["Gestão de conteúdo", "Política comercial"], "Campanha que chega atrasada não vende."),
    "Insumo de trabalho": (["Locação / venda", "Estoque", "Acervo", "Garantias"], ["Manutenção", "SLA", "Insegurança técnica"], "Ao vivo não tem segunda chance."),
    "Experiência": (["Customização", "Solução 360", "Prova social"], ["Preço", "Percepção de valor"], "Espaço sem experiência é só metro quadrado."),
    "Entretenimento": (["Customização", "Acervo", "Locação / venda"], ["Consumo de energia", "Manutenção", "Momento financeiro"], "Quem encanta, faz voltar."),
    "Status": (["Liderança", "Premiações", "Pioneira", "Customização"], ["Preço", "Concorrência"], "Seja a referência que os outros vão copiar."),
}
wm = sheet("Mapa de Venda", MAGENTA)
title(wm, "Mapa de Venda: propósito → diferencial → resultado", "Insight do time comercial. Descubra o propósito de compra e conduza a conversa pela linha dele.", 9)
cols = ["Propósito de compra", "O que o cliente quer", "Pergunta para identificar", "Diferenciais que provam", "Resultados que entrega",
        "Objeções típicas", "Cases", "Frase de efeito", "Nichos onde aparece"]
w = [20, 38, 36, 30, 30, 30, 22, 30, 36]
header(wm, 3, cols, w)
for i, (prop, quer, perg, res) in enumerate(PROPOSITOS):
    difs, objs, frase = MAPA_EXTRA[prop]
    cases = "; ".join(c[0] for c in CASES if prop in c[2])
    nichos = "; ".join(k for k, v in NICHO_MAPA.items() if prop in v[0])
    row(wm, 4 + i, [prop, quer, perg, "; ".join(difs), res, "; ".join(objs), cases, frase, nichos], w, zebra=(i % 2 == 1))
    wm.cell(row=4 + i, column=8).font = Font(name=F, size=10, bold=True, color=MAGENTA)
wm.freeze_panes = "B4"

# ------------------------------------------------------------------ Diferenciais
wd = sheet("Diferenciais")
title(wd, "Diferenciais da LedWave", "Use 3 a 5 por apresentação, escolhidos pelo propósito de compra. 'Confirmar' = validar com o comercial antes de usar com cliente.", 5)
w = [26, 66, 40, 20, 28]
header(wd, 3, ["Diferencial", "Como falar", "Prova / o que levar", "Status", "Observação"], w)
for i, d in enumerate(DIFERENCIAIS):
    row(wd, 4 + i, list(d) + [""], w, edit_cols=(4, 5))
ld = 3 + len(DIFERENCIAIS)
wd.conditional_formatting.add(f"D4:D{ld}", FormulaRule(formula=['LEFT($D4,9)="Confirmar"'], font=Font(color="B00020", bold=True)))
wd.conditional_formatting.add(f"D4:D{ld}", FormulaRule(formula=['$D4="OK"'], font=Font(color="1B7F3B", bold=True)))
wd.freeze_panes = "B4"

# ------------------------------------------------------------------ Objeções
wo = sheet("Objeções")
title(wo, "Objeções: resposta curta + prova", "Método: ouvir até o fim → acolher → explorar ('além disso, tem mais algo?') → responder com prova → confirmar.", 6)
w = [16, 30, 70, 32, 16, 34]
header(wo, 3, ["Categoria", "Objeção", "Resposta", "Leve / prova", "Funcionou?", "Variação que funcionou"], w)
objs = [(cat, o, resp, prova) for o, cat, resp, prova in OBJECOES_COMERCIAL]
objs += [("Decisão", "Vou pensar.", OBJECOES_GERAIS[1][2], "Pergunta de isolamento"),
         ("Percepção", "Já tenho TV / monitor.", OBJECOES_GERAIS[2][2], "Demonstração lado a lado"),
         ("Decisão", "Preciso falar com meu sócio / diretoria.", OBJECOES_GERAIS[7][2], "Resumo executivo de 1 página")]
objs += [("Nicho: " + n["nicho"], n["objecao"], n["resposta"], "") for n in NICHOS]
for i, o in enumerate(objs):
    row(wo, 4 + i, list(o) + ["", ""], w, edit_cols=(5, 6))
lo = 3 + len(objs)
dv = DataValidation(type="list", formula1='"Sim,Parcial,Não"', allow_blank=True)
wo.add_data_validation(dv); dv.add(f"E4:E{lo}")
wo.freeze_panes = "C4"
wo.auto_filter.ref = f"A3:F{lo}"

# ------------------------------------------------------------------ Cases
wc = sheet("Cases")
title(wc, "Cases de referência", "Preencha desafio, entrega e resultado: é isso que entra nas apresentações. Só use com cliente após autorização.", 8)
w = [22, 30, 30, 34, 34, 30, 16, 22]
header(wc, 3, ["Case", "Nicho", "Propósitos de compra", "Desafio do cliente", "O que a LedWave entregou", "Resultado", "Autorizado?", "Status"], w)
for i, (case, seg, props, st) in enumerate(CASES):
    row(wc, 4 + i, [case, seg, props, "", "", "", "", st], w, edit_cols=(4, 5, 6, 7, 8))
    wc.row_dimensions[4 + i].height = 40
for rr in range(4 + len(CASES), 14 + len(CASES)):
    for cc in range(1, 9):
        cell = wc.cell(row=rr, column=cc)
        cell.fill, cell.border = PatternFill("solid", fgColor=EDIT), BORDER
dv = DataValidation(type="list", formula1='"Sim,Não,A confirmar"', allow_blank=True)
wc.add_data_validation(dv); dv.add(f"G4:G{13 + len(CASES)}")
wc.freeze_panes = "B4"

# ------------------------------------------------------------------ Descoberta
wq = sheet("Descoberta")
title(wq, "Descoberta e orçamento por nicho", "Perguntas SPIN (faça poucas de Situação, mais de Implicação) e o que levantar para orçar certo.", 7)
w = [24, 36, 36, 36, 36, 42, 34]
header(wq, 3, ["Nicho", "S · Situação", "P · Problema", "I · Implicação", "N · Necessidade", "Dados para orçar", "Armadilha de orçamento"], w)
for i, n in enumerate(NICHOS):
    s, p, im, ne = SPIN[n["nicho"]]
    tec, qual, arm = ORCAMENTO[n["nicho"]]
    row(wq, 4 + i, [n["nicho"], s, p, im, ne, f"{tec}\nQualificação: {qual}", arm], w, zebra=(i % 2 == 1))
wq.freeze_panes = "B4"

# ------------------------------------------------------------------ Treino
wt = sheet("Treino")
title(wt, "Treino rápido", "O essencial de apresentação, escuta e prospecção. Roleplay e técnicas completas estão na versão completa.", 3)
w = [30, 70, 50]
for i, x in enumerate(w, start=1):
    wt.column_dimensions[get_column_letter(i)].width = x
blocos_t = [
    ("Apresentação", [
        ("Comece pelo porquê", "O propósito do cliente antes do painel.", "'Sua vitrine pode vender 24h' antes de 'painel P3,9'."),
        ("Cliente herói, LedWave guia", "O cliente é o protagonista; a LedWave já fez isso antes.", "'Você quer ser a escolhida da rua. Nós ajudamos centenas a chegar lá.'"),
        ("Regra de 3", "Três provas, três ganhos, três números.", "Receita, experiência, operação."),
        ("Frase + pausa", "Frase curta, depois 3 segundos de silêncio.", "'Vitrine parada é vendedor dormindo.'"),
        ("Termine no ponto alto", "Nunca 'alguma dúvida?'. Termine com duas datas.", "'Visita técnica na quinta ou na sexta?'"),
    ]),
    ("Escuta ativa", [
        ("Regra 70/30", "Na descoberta, o cliente fala 70% do tempo.", "Se você fala mais, está vendendo cedo demais."),
        ("Silêncio de 3 s", "Espere depois da resposta: vem a informação mais valiosa.", "—"),
        ("Espelhamento", "Repita as últimas palavras em tom de pergunta.", "'...é um caos.' → 'Um caos?'"),
        ("Paráfrase", "Resuma com as palavras dele e confirme.", "'Se entendi bem, o principal é X. Correto?'"),
        ("'O que mais?'", "A segunda dor costuma ser a mais importante.", "—"),
    ]),
    ("Prospecção (cadência de 14 dias)", [
        ("D0 · Pesquisa", "Ache um gatilho: inauguração, reforma, lançamento, feira, troca de liderança.", "Instagram, Google Maps, LinkedIn, imprensa."),
        ("D1 · LinkedIn", "Conexão com nota personalizada, sem pitch.", "'Vi {gatilho}. Trabalho com comunicação digital para {nicho}.'"),
        ("D2 · Ligação + WhatsApp", "Objetivo: agendar 20 min, não vender.", "Permissão → motivo → pergunta de dor → mini-prova → duas datas."),
        ("D4 · E-mail", "Um insight do nicho + uma imagem antes/depois. Máx. 90 palavras.", "Assunto: {Empresa} + uma ideia para {dor}."),
        ("D6 a D11", "Nova ligação, case do nicho, vídeo de 45 s com mockup do espaço.", "Quebre o padrão."),
        ("D14 · Encerramento", "E-mail educado de término.", "'Encerro por aqui ou volto a falar em {mês}?'"),
    ]),
]
r = 4
for nome, itens in blocos_t:
    section(wt, r, nome)
    header(wt, r + 1, ["Técnica / etapa", "Como fazer", "Exemplo"], color=NAVY)
    for j, it in enumerate(itens):
        row(wt, r + 2 + j, list(it), w)
    r += 3 + len(itens)

# ------------------------------------------------------------------ Ideias
wi = sheet("Ideias", MAGENTA)
title(wi, "Banco de Ideias dos Gestores", "Uma linha por ideia. A linha 4 é um exemplo. 'Aprovada' = levar para a aba Nichos ou Mapa de Venda.", 8)
w = [12, 18, 26, 22, 60, 30, 20, 10]
header(wi, 3, ["Data", "Gestor", "Nicho", "Tipo", "Ideia (texto pronto para usar)", "Testada? Resultado", "Status", "Votos"], w)
ex = [date(2026, 10, 19), "Exemplo (apagar)", "Farmácias e Drogarias", "Pergunta de abertura",
      "Se o seu cliente estivesse com dor de cabeça na esquina, ele veria a sua farmácia antes da concorrente?",
      "2 reuniões: abriu conversa sobre fachada", "Em teste em campo", 3]
row(wi, 4, ex, w, bold_first=False)
for rr in range(4, 305):
    for cc in range(1, 9):
        cell = wi.cell(row=rr, column=cc)
        cell.border, cell.alignment = BORDER, WRAP
        if rr > 4:
            cell.fill, cell.font = PatternFill("solid", fgColor=EDIT), Font(name=F, size=10)
    wi.cell(row=rr, column=1).number_format = "DD/MM/YYYY"
tipos = ["Pergunta de abertura", "Argumento", "Frase de efeito", "Objeção + resposta", "Case", "Diferencial", "Propósito de compra", "Pergunta de descoberta"]
for colL, ref in (("C", "=Listas!$A$2:$A$60"), ("D", "=Listas!$B$2:$B$20"), ("G", "=Listas!$C$2:$C$10")):
    dv = DataValidation(type="list", formula1=ref, allow_blank=True)
    wi.add_data_validation(dv); dv.add(f"{colL}4:{colL}304")
wi.conditional_formatting.add("A4:H304", FormulaRule(formula=['$G4="Aprovada (levar à matriz)"'], fill=PatternFill("solid", fgColor="D7F2DF")))
wi.conditional_formatting.add("A4:H304", FormulaRule(formula=['$G4="Descartada"'], font=Font(color="999999", strike=True)))
wi.freeze_panes = "A4"
wi.auto_filter.ref = "A3:H304"
wi.column_dimensions["J"].width, wi.column_dimensions["K"].width = 26, 10
wi["J3"] = "Resumo"
wi["J3"].font = Font(name=F, size=11, bold=True, color=NAVY)
for i, s in enumerate(STATUS_IDEIA):
    wi.cell(row=4 + i, column=10, value=s).font = Font(name=F, size=10)
    wi.cell(row=4 + i, column=11, value=f"=COUNTIF($G$4:$G$304,J{4+i})").font = Font(name=F, size=10)

# ------------------------------------------------------------------ Listas
wl = sheet("Listas", "999999")
for i, h in enumerate(["Nichos", "Tipos", "Status"], start=1):
    c = wl.cell(row=1, column=i, value=h)
    c.font, c.fill = Font(name=F, size=10, bold=True, color=WHITE), PatternFill("solid", fgColor=NAVY)
    wl.column_dimensions[get_column_letter(i)].width = 34
for i in range(59):
    wl.cell(row=2 + i, column=1, value=f'=IF(Nichos!A{4+i}="","",Nichos!A{4+i})')
for i, t in enumerate(tipos):
    wl.cell(row=2 + i, column=2, value=t)
for i, t in enumerate(STATUS_IDEIA):
    wl.cell(row=2 + i, column=3, value=t)
wl.sheet_state = "hidden"

wb.active = 0
wb.save(OUT)
print("ok", OUT)
